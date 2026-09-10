"""One-off, target-checked 10 September 2026 payroll release.

Run only through the MAM project credential helper after release/import approval.
prepare <source-directory> <private-json-output> performs read-only reconciliation.
migrate <reviewed-filename> <sha256> applies only one of the two named migrations.
import <private-json-output> stores all 28 records and syncs exact active matches.
verify reads counts only. Keep all payroll inputs and outputs outside Git.
"""
import hashlib,json,os,pathlib,re,sys,urllib.request
ROOT=pathlib.Path(__file__).resolve().parents[1]
REF='xtwlrjqbusadghavchqm'; ORG='fwlmbmospuwoaosffvmw'
def api(path,body=None):
    req=urllib.request.Request('https://api.supabase.com/v1/projects/'+REF+path,headers={'Authorization':'Bearer '+os.environ['SUPABASE_ACCESS_TOKEN'],'Content-Type':'application/json','User-Agent':'supabase-cli/2.54.11'},data=None if body is None else json.dumps(body).encode())
    with urllib.request.urlopen(req,timeout=90) as r:return json.load(r)
def query(sql):return api('/database/query',{'query':sql})
project=api('');assert project['id']==REF and project['organization_id']==ORG
def literal(value):return "'"+str(value).replace("'","''")+"'"
def normal(name):return re.sub('[^a-z0-9]','',name.casefold())
if sys.argv[1]=='prepare':
    import openpyxl
    roster=query("select p.id,p.full_name,p.company_key,p.is_active,d.first_name,d.last_name from public.internal_org_people p left join private.employee_directory_details d on d.person_id=p.id where p.node_kind='person'")
    entries=[]
    directory=pathlib.Path(sys.argv[2])
    for company,filename in [('mam','Mobility_Access_Modifications_Pty_Ltd_-_Leave_Balances.xlsx'),('equipsy','Equipsy_Pty_Ltd_-_Leave_Balances.xlsx')]:
        file=directory/filename; sha=hashlib.sha256(file.read_bytes()).hexdigest(); rows={}
        workbook=openpyxl.load_workbook(file,read_only=True,data_only=True)
        for sheet in workbook:
            assert sheet.cell(3,1).value=='As at 10 September 2026'
            kind={'Annual Leave':'annual_hours','Sick Leave':'sick_hours'}[sheet.title.strip()]
            for name,leave_type,hours in sheet.iter_rows(min_row=6,values_only=True):
                if not name or not leave_type:continue
                assert isinstance(hours,(int,float)) and -10000<=hours<=100000
                row=rows.setdefault(name,{'company_key':company,'employee_name':name,'as_of':'2026-09-10','source_file':filename,'source_sha256':sha})
                assert kind not in row;row[kind]=hours
        for row in rows.values():
            assert 'annual_hours' in row and 'sick_hours' in row
            matches=[p for p in roster if normal(row['employee_name']) in {normal(p['full_name']),normal((p['first_name'] or '')+' '+(p['last_name'] or ''))}]
            active=[p for p in matches if p['company_key']==company and p['is_active']]
            inactive=[p for p in matches if p['company_key']==company and not p['is_active']]
            person=active[0] if len(active)==1 else inactive[0] if len(active)==0 and len(inactive)==1 else None
            row['person_id']=None if person is None else person['id']
            row['match_status']='linked' if person and person['is_active'] else 'inactive' if person else 'company_mismatch' if matches and all(p['company_key']!=company for p in matches) else 'unmatched'
            entries.append(row)
    assert len(entries)==28
    out=pathlib.Path(sys.argv[3]);out.write_text(json.dumps(entries));out.chmod(0o600)
    print(json.dumps({'verifiedProject':REF,'records':len(entries),'balances':2*len(entries),'statuses':{s:sum(e['match_status']==s for e in entries) for s in ['linked','inactive','company_mismatch','unmatched']}}))
elif sys.argv[1]=='migrate':
    filename=sys.argv[2];assert filename in ['20260910012953_payroll_balance_visibility.sql','20260910015007_payroll_import_register.sql']
    sql=(ROOT/'supabase/migrations'/filename).read_text(); assert hashlib.sha256(sql.encode()).hexdigest()==sys.argv[3]
    version,name=filename[:-4].split('_',1)
    existing=query('select version from supabase_migrations.schema_migrations where version='+literal(version))
    if existing:print('Already applied '+version);sys.exit(0)
    assert sql.count('\ncommit;')==1
    history='insert into supabase_migrations.schema_migrations(version,name,statements) values('+literal(version)+','+literal(name)+',ARRAY['+literal(sql)+']);'
    sql=sql.replace('begin;','begin;\nset local lock_timeout=\'5s\';\nset local statement_timeout=\'75s\';',1).replace('\ncommit;','\n'+history+"\nnotify pgrst,'reload schema';\ncommit;")
    query(sql);print('Applied '+version)
elif sys.argv[1]=='import':
    entries=json.loads(pathlib.Path(sys.argv[2]).read_text());assert len(entries)==28
    sql="""begin;
set local lock_timeout='5s';set local statement_timeout='75s';
create temp table task_payroll_rows on commit drop as
select * from jsonb_to_recordset(PAYLOAD::jsonb) as x(company_key text,employee_name text,annual_hours numeric,sick_hours numeric,as_of date,source_file text,source_sha256 text,person_id uuid,match_status text);
do $$begin
  if exists(select 1 from task_payroll_rows s left join public.internal_org_people p on p.id=s.person_id
    where s.person_id is not null and (p.id is null or p.company_key<>s.company_key or p.is_active<>(s.match_status='linked')
      or regexp_replace(lower(p.full_name),'[^a-z0-9]','','g')<>regexp_replace(lower(s.employee_name),'[^a-z0-9]','','g'))) then
    raise exception 'Employee mapping changed or is not an exact normalized same-company name';end if;
  if exists(select person_id from task_payroll_rows where match_status='linked' group by person_id having count(*)>1) then
    raise exception 'Duplicate employee mapping';end if;
end$$;
insert into private.payroll_balance_imports(company_key,employee_name,annual_hours,sick_hours,as_of,source_file,source_sha256,person_id,match_status)
select company_key,employee_name,annual_hours,sick_hours,as_of,source_file,source_sha256,person_id,match_status from task_payroll_rows
on conflict(source_sha256,company_key,employee_name,as_of) do nothing;
do $$begin
  if (select count(*) from private.payroll_balance_imports i join task_payroll_rows s using(source_sha256,company_key,employee_name,as_of)
    where i.annual_hours=s.annual_hours and i.sick_hours=s.sick_hours and i.person_id is not distinct from s.person_id and i.match_status=s.match_status)<>28 then
    raise exception 'Import reconciliation failed';end if;
end$$;
insert into private.employee_annual_leave_balances as b(person_id,hours,as_of,is_example,updated_by,source_import_id)
select i.person_id,i.annual_hours,i.as_of,false,null,i.id from private.payroll_balance_imports i
join task_payroll_rows s using(source_sha256,company_key,employee_name,as_of) where i.match_status='linked'
on conflict(person_id) do update set hours=excluded.hours,as_of=excluded.as_of,is_example=false,updated_by=null,source_import_id=excluded.source_import_id,updated_at=now()
where b.as_of<=excluded.as_of and b.source_import_id is distinct from excluded.source_import_id;
insert into private.employee_sick_leave_balances as b(person_id,hours,as_of,updated_by,source_import_id)
select i.person_id,i.sick_hours,i.as_of,null,i.id from private.payroll_balance_imports i
join task_payroll_rows s using(source_sha256,company_key,employee_name,as_of) where i.match_status='linked'
on conflict(person_id) do update set hours=excluded.hours,as_of=excluded.as_of,updated_by=null,source_import_id=excluded.source_import_id,updated_at=now()
where b.as_of<=excluded.as_of and b.source_import_id is distinct from excluded.source_import_id;
insert into private.annual_leave_balance_history(person_id,hours,as_of,is_example,actor_user_id,source_import_id)
select b.person_id,b.hours,b.as_of,false,null,b.source_import_id from private.employee_annual_leave_balances b
join private.payroll_balance_imports i on i.id=b.source_import_id join task_payroll_rows s on s.source_sha256=i.source_sha256 and s.company_key=i.company_key and s.employee_name=i.employee_name and s.as_of=i.as_of
on conflict do nothing;
insert into private.sick_leave_balance_history(person_id,hours,as_of,actor_user_id,source_import_id)
select b.person_id,b.hours,b.as_of,null,b.source_import_id from private.employee_sick_leave_balances b
join private.payroll_balance_imports i on i.id=b.source_import_id join task_payroll_rows s on s.source_sha256=i.source_sha256 and s.company_key=i.company_key and s.employee_name=i.employee_name and s.as_of=i.as_of
on conflict do nothing;
commit;
""".replace('PAYLOAD',literal(json.dumps(entries)))
    query(sql);print('Imported all 28 payroll records and synchronized unambiguous active matches.')
elif sys.argv[1]=='verify':
    print(json.dumps(query("""select i.match_status,count(*) records,count(*) filter(where b.source_import_id=i.id and b.hours=i.annual_hours and b.as_of=i.as_of) annual_synced,count(*) filter(where s.source_import_id=i.id and s.hours=i.sick_hours and s.as_of=i.as_of) sick_synced from private.payroll_balance_imports i left join private.employee_annual_leave_balances b on b.person_id=i.person_id left join private.employee_sick_leave_balances s on s.person_id=i.person_id where i.as_of='2026-09-10' group by i.match_status order by i.match_status""")))
    print(json.dumps(query("""select version,name from supabase_migrations.schema_migrations where version in ('20260910012953','20260910015007') order by version""")))
